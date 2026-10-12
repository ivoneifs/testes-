"""Converte arquivos enviados (PDF, imagem, DOC/DOCX, XLS/XLSX, PPTX, CSV/TXT/JSON) em itens
de conteúdo da OpenAI Responses API.

PDF e imagem vão como arquivo/imagem (o modelo lê direto). Word e Excel viram
texto aqui no servidor — tabelas e planilhas linha a linha, células separadas
por " | " — porque é assim que os escores costumam vir.
"""
from __future__ import annotations

import base64
import io
import json
import re
import zipfile

MAX_TEXT = 60_000  # caracteres por arquivo

ACCEPTED_EXT = ('.pdf', '.png', '.jpg', '.jpeg', '.webp', '.gif',
                '.doc', '.docx', '.xls', '.xlsx', '.pptx', '.txt', '.csv', '.md', '.json')
FORMATS_LABEL = 'PDF, JPG, PNG, DOC, DOCX, XLS, XLSX, PPTX, CSV, TXT ou JSON'


def accepted(filename: str, mime: str = '') -> bool:
    name = (filename or '').lower()
    return name.endswith(ACCEPTED_EXT) or (mime or '').startswith(('image/', 'text/')) or mime == 'application/pdf'


def _cell(v) -> str:
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _rows_text(rows) -> str:
    out = []
    for row in rows:
        vals = [_cell(v) for v in row]
        while vals and not vals[-1]:
            vals.pop()
        if any(vals):
            out.append(' | '.join(vals))
    return '\n'.join(out)


def docx_text(data: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for i, t in enumerate(doc.tables, 1):
        parts.append(f'[Tabela {i}]')
        parts.append(_rows_text([c.text for c in r.cells] for r in t.rows))
    return '\n'.join(parts)


def xlsx_text(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        body = _rows_text(ws.iter_rows(values_only=True))
        if body:
            parts.append(f'[Planilha: {ws.title}]\n{body}')
    return '\n\n'.join(parts)


def xls_text(data: bytes) -> str:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    parts = []
    for sh in wb.sheets():
        body = _rows_text(sh.row_values(r) for r in range(sh.nrows))
        if body:
            parts.append(f'[Planilha: {sh.name}]\n{body}')
    return '\n\n'.join(parts)


def pptx_text(data: bytes) -> str:
    """Texto de cada slide (e das anotações), sem dependência: lê o XML do .pptx."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        def num(n):
            m = re.search(r'(\d+)\.xml$', n)
            return int(m.group(1)) if m else 0
        def runs(n):
            xml = z.read(n).decode('utf-8', 'ignore')
            paras = re.findall(r'<a:p\b.*?</a:p>', xml, re.S)
            out = [''.join(re.findall(r'<a:t[^>]*>(.*?)</a:t>', p, re.S)) for p in paras]
            return '\n'.join(_unxml(t) for t in out if t.strip())
        slides = sorted((n for n in z.namelist() if re.match(r'ppt/slides/slide\d+\.xml$', n)), key=num)
        notes = {num(n): n for n in z.namelist() if re.match(r'ppt/notesSlides/notesSlide\d+\.xml$', n)}
        parts = []
        for n in slides:
            body = runs(n)
            if num(n) in notes:
                nt = runs(notes[num(n)])
                body += f'\n[Anotações] {nt}' if nt else ''
            if body.strip():
                parts.append(f'[Slide {num(n)}]\n{body}')
        return '\n\n'.join(parts)


def _unxml(t: str) -> str:
    return (t.replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"')
             .replace('&apos;', "'").replace('&amp;', '&'))


def json_text(data: bytes) -> str:
    raw = data.decode('utf-8-sig', 'replace')
    try:
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=1)
    except ValueError:
        return raw


def doc_text(data: bytes) -> str:
    """.doc (Word 97-2003) sem dependência: recupera os trechos de texto
    legíveis (UTF-16 e 8-bit). Perde formatação, mas mantém os números."""
    # Word guarda o texto em UTF-16 ou, quando "comprimido", em cp1252: usa o que render mais.
    t16 = '\n'.join(r.decode('utf-16le', 'ignore')
                    for r in re.findall(rb'(?:[\x20-\x7e\xa0-\xff][\x00]){4,}', data))
    t8 = '\n'.join(r.decode('cp1252', 'ignore')
                   for r in re.findall(rb'[\x20-\x7e\xa0-\xff\r\n\t]{6,}', data))
    return re.sub(r'\n{3,}', '\n\n', max(t16, t8, key=len))


def to_content(filename: str, mime: str, data: bytes) -> list[dict]:
    """Itens de conteúdo para um arquivo. Levanta ValueError se não suportado."""
    name = (filename or 'arquivo').lower()
    mime = mime or ''
    b64 = lambda: base64.b64encode(data).decode('ascii')
    if mime == 'application/pdf' or name.endswith('.pdf'):
        return [{'type': 'input_file', 'filename': filename,
                 'file_data': f'data:application/pdf;base64,{b64()}'}]
    if mime.startswith('image/') or name.endswith(('.png', '.jpg', '.jpeg', '.webp', '.gif')):
        if not mime.startswith('image/'):
            mime = 'image/jpeg' if name.endswith(('.jpg', '.jpeg')) else f'image/{name.rsplit(".", 1)[-1]}'
        return [{'type': 'input_image', 'image_url': f'data:{mime};base64,{b64()}', 'detail': 'high'}]
    if name.endswith('.docx'):
        text = docx_text(data)
    elif name.endswith('.xlsx'):
        text = xlsx_text(data)
    elif name.endswith('.xls'):
        text = xls_text(data)
    elif name.endswith('.doc'):
        text = doc_text(data)
    elif name.endswith('.pptx'):
        text = pptx_text(data)
    elif name.endswith('.json') or mime == 'application/json':
        text = json_text(data)
    elif mime.startswith('text/') or name.endswith(('.txt', '.csv', '.md')):
        text = data.decode('utf-8-sig', 'replace')
    else:
        raise ValueError(f'{filename}: formato não suportado (use {FORMATS_LABEL})')
    text = text.strip()
    if not text:
        raise ValueError(f'{filename}: não foi possível ler texto do arquivo')
    return [{'type': 'input_text', 'text': f'--- Arquivo: {filename} ---\n{text[:MAX_TEXT]}'}]
